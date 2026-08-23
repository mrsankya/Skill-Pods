import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette
    PRIMARY_COLOR = RGBColor(38, 21, 67)     # #261543 Deep Purple
    SECONDARY_COLOR = RGBColor(109, 78, 199) # #6d4ec7 Neon Lavender Purple
    DARK_TEXT = RGBColor(30, 41, 59)         # Slate 800
    MUTED_TEXT = RGBColor(100, 116, 139)     # Slate 500
    EMERALD = RGBColor(5, 150, 105)          # Emerald 600

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("SKILLPODS — TECHNICAL APPROACH & ARCHITECTURE SPECIFICATION")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("Smart India Hackathon (SIH) Technical Submission & Architecture Documentation")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper for Headings
    def add_custom_heading(text, level=1):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(14)
            run.font.color.rgb = PRIMARY_COLOR
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = SECONDARY_COLOR
        else:
            run.font.size = Pt(11)
            run.font.color.rgb = DARK_TEXT
        return h

    # Section 1: Executive Overview
    add_custom_heading("1. Executive Summary & Core Value Proposition", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "SkillPods is an enterprise-grade collaborative platform designed to bridge the gap between Small & Medium Enterprises (SMEs), "
        "Engineering Students, Industry Mentors, and Universities. The platform transforms real-world SME business challenges into production-ready "
        "commercial software built by agile student Skill Pods, guided by verified industry architects, and protected through cryptographic skill verification "
        "and automated milestone escrow smart gating."
    )

    # Section 2: High-Level Architecture
    add_custom_heading("2. High-Level 4-Tier System Architecture", 1)
    
    arch_table = doc.add_table(rows=5, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    arch_headers = ["Layer / Tier", "Core Technologies", "Technical Responsibilities"]
    
    for i, head in enumerate(arch_headers):
        cell = arch_table.cell(0, i)
        cell.text = head
        set_cell_background(cell, "261543")
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)
        set_cell_margins(cell, 120, 120, 150, 150)

    arch_data = [
        ("1. Presentation Tier (Frontend)", "React 19, TypeScript, Tailwind CSS v4, Motion v12, Canvas 3D Avatars", "Role-adaptive responsive dashboards for Students, SMEs, Mentors, and Colleges with low-latency glassmorphism rendering."),
        ("2. API & Gateway Tier", "Node.js 20+, Express.js, Vite Middleware Integration, Zod Validation", "Secure REST API routing, session authorization, payload validation, CORS management, and real-time event streaming."),
        ("3. Intelligence & Service Tier", "Google Gemini API (@google/genai), pgvector Cosine Matcher, SHA-256 Engine", "Automated skill matching, pod fit scoring, anomaly telemetry radar, and cryptographic credential passport minting."),
        ("4. Data & Integration Tier", "PostgreSQL, TimescaleDB, GitHub Webhook Ingest, Razorpay Escrow Highway", "Relational persistence, time-series commit and code coverage metrics, automated milestone release escrow smart highway.")
    ]

    for row_idx, data in enumerate(arch_data, start=1):
        for col_idx, text in enumerate(data):
            cell = arch_table.cell(row_idx, col_idx)
            cell.text = text
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8F7FC")
            else:
                set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, 100, 100, 120, 120)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3: Core Algorithms & Formulations
    add_custom_heading("3. Core Algorithms & Mathematical Formulations", 1)

    add_custom_heading("A. Multi-Factor AI Pod Fit Score Engine", 2)
    p = doc.add_paragraph()
    p.add_run(
        "To objectively determine student and pod compatibility with industry problems, SkillPods utilizes a weighted vector scoring formulation:"
    )
    
    formula_p = doc.add_paragraph()
    formula_p.paragraph_format.left_indent = Inches(0.3)
    formula_run = formula_p.add_run("MatchScore(P, S) = (0.40 × S_semantic) + (0.25 × S_velocity) + (0.20 × S_mentor) + (0.15 × S_coverage)")
    formula_run.font.bold = True
    formula_run.font.color.rgb = PRIMARY_COLOR

    p_factors = doc.add_paragraph()
    p_factors.add_run("• S_semantic (40%): ").bold = True
    p_factors.add_run("Cosine similarity between student competency embedding vectors and problem requirement vectors.\n")
    p_factors.add_run("• S_velocity (25%): ").bold = True
    p_factors.add_run("Historical sprint completion velocity benchmarked against cohort baselines (e.g. 94/100).\n")
    p_factors.add_run("• S_mentor (20%): ").bold = True
    p_factors.add_run("Seniority and availability score of the assigned domain architect (e.g. 4.98 ★).\n")
    p_factors.add_run("• S_coverage (15%): ").bold = True
    p_factors.add_run("Empirical automated test coverage achieved across past delivered modules (e.g. 94.2% Cypress E2E).")

    add_custom_heading("B. Cryptographic Skill Passport Integrity Formula", 2)
    p_pass = doc.add_paragraph()
    p_pass.add_run(
        "Every student credential is cryptographically stamped into an immutable digital passport to prevent resume fraud and ensure NAAC/NBA compliance:\n"
    )
    formula_p2 = doc.add_paragraph()
    formula_p2.paragraph_format.left_indent = Inches(0.3)
    formula_run2 = formula_p2.add_run("PassportSignature = SHA-256( RollNo || MentorPubKey || VerifiedSkills || MergedPRs || Timestamp )")
    formula_run2.font.bold = True
    formula_run2.font.color.rgb = PRIMARY_COLOR

    add_custom_heading("C. Pod Health & Bottleneck Radar Formula", 2)
    p_radar = doc.add_paragraph()
    p_radar.add_run("The platform monitors real-time GitHub telemetry to compute a continuous Pod Health Score (H):\n")
    formula_p3 = doc.add_paragraph()
    formula_p3.paragraph_format.left_indent = Inches(0.3)
    formula_run3 = formula_p3.add_run("H = 100 - [ (15 × BlockedPRs) + (20 × VelocityDropPct) + (10 × WorkloadGiniImbalance) ]")
    formula_run3.font.bold = True
    formula_run3.font.color.rgb = PRIMARY_COLOR

    # Section 4: Two-Phase Milestone Escrow Protocol
    add_custom_heading("4. Two-Phase Milestone Escrow Gating Protocol", 1)
    p_escrow = doc.add_paragraph()
    p_escrow.add_run(
        "To protect both SME capital and student efforts, the platform enforces an automated two-phase release gate:\n"
    )
    p_steps = doc.add_paragraph()
    p_steps.add_run("1. SME Funding Lock: ").bold = True
    p_steps.add_run("SME deposits project milestone bounty into the SkillPods Escrow Vault prior to sprint initiation.\n")
    p_steps.add_run("2. Sprint Delivery & Testing: ").bold = True
    p_steps.add_run("Student Pod develops features and submits pull requests backed by ≥90% Cypress/Jest test suites.\n")
    p_steps.add_run("3. Phase 1 — Mentor Architecture Gate: ").bold = True
    p_steps.add_run("Industry Mentor inspects code architecture, security, and test pass rates. Only upon explicit mentor sign-off is the next sprint unlocked.\n")
    p_steps.add_run("4. Phase 2 — SME Commercial Acceptance: ").bold = True
    p_steps.add_run("SME validates the live deployed staging environment against API latency SLAs (sub-50ms) and triggers automated escrow payout disbursement.")

    # Section 5: Role-by-Role Feature Matrix
    add_custom_heading("5. Role-Specific Feature Matrix", 1)

    feat_table = doc.add_table(rows=5, cols=3)
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    feat_headers = ["Role Dashboard", "Key Modules Implemented", "Business / Academic Impact"]

    for i, head in enumerate(feat_headers):
        cell = feat_table.cell(0, i)
        cell.text = head
        set_cell_background(cell, "261543")
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)
        set_cell_margins(cell, 120, 120, 150, 150)

    feat_data = [
        ("Student Dashboard", "• AI Skill Matcher\n• Project Marketplace\n• Cryptographic Passport\n• Earnings Wallet & Payouts\n• Enterprise v2 Upgrade", "Enables students to monetize college capstone code, receive mentor verified credentials, and earn industry stipends."),
        ("SME Dashboard", "• AI Pod Recommendations\n• Pod Side-by-Side Comparison\n• Student Solution Buyer Portal\n• Milestone Acceptance Gate\n• Codebase Reuse Engine", "Reduces corporate software development costs by 60% and accelerates time-to-market using pre-built student foundations."),
        ("Mentor Dashboard", "• Sprint Milestone Gating\n• Skill Verification Stamps\n• Contribution Score Auditing\n• Pod Health Radar & Alerts\n• Permanent Audit History", "Empowers industry staff engineers to enforce production architecture standards and mentor next-gen talent efficiently."),
        ("College Dashboard", "• Student Skill Passport Directory\n• Innovation & IP Registry\n• Inter-Department Analytics\n• Placement Readiness Index\n• NAAC/NBA Accreditation Reports", "Provides Deans and Placement Cells with verifiable institutional data on live industry projects and student employability.")
    ]

    for row_idx, data in enumerate(feat_data, start=1):
        for col_idx, text in enumerate(data):
            cell = feat_table.cell(row_idx, col_idx)
            cell.text = text
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8F7FC")
            else:
                set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, 100, 100, 120, 120)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 6: SIH Presentation Deck Slide-by-Slide Guide
    add_custom_heading("6. Recommended SIH Presentation Deck Mapping", 1)
    
    p_slides = doc.add_paragraph()
    p_slides.add_run("Slide 1 — System Architecture: ").bold = True
    p_slides.add_run("Present the 4-Tier diagram (Presentation, API Gateway, Intelligence Engines, Persistence Layer).\n")
    p_slides.add_run("Slide 2 — Technical Innovations & AI: ").bold = True
    p_slides.add_run("Highlight the Multi-Factor Pod Fit Score formulation and SHA-256 Cryptographic Passport minting.\n")
    p_slides.add_run("Slide 3 — Trust & Security Protocol: ").bold = True
    p_slides.add_run("Walk through the Two-Phase Milestone Escrow Gating Protocol demonstrating fraud prevention.\n")
    p_slides.add_run("Slide 4 — Market Impact & Multi-Role Cockpit: ").bold = True
    p_slides.add_run("Showcase the monetization models (Commercial License, IP Buyout, Upgrades) and College IP Registry.")

    # Save Document
    filename = "SkillPods_SIH_Technical_Approach.docx"
    doc.save(filename)
    print(f"Successfully generated {filename}")

if __name__ == '__main__':
    create_document()
